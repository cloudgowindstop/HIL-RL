#!/usr/bin/env python3
"""Plan and download BOS HDF5 batches from the board or 9-device inventory.

The board is treated as source data, not repaired or inferred. Missing station
and task values become ``unknown_station`` and ``unknown_task``. BOS path
numbers never participate in metadata resolution.

``plan-incremental`` compares exact BOS roots against historical successful
results and writes repaired/new batches into an isolated staging directory.

``inventory`` lists BOS batch roots and compares them with local downloads.
It never calls ``bos sync``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import h5py
from openpyxl import Workbook, load_workbook


DEFAULT_BCECMD = Path("/media/linux-bcecmd-0.5.1/bcecmd")
DEFAULT_BOS_BASE = "bos:/bd-dp-ten-6spt6-scjd"
DEFAULT_OUTPUT_ROOT = Path("/media/jushen/project-rl-dataset/raw_data_0831download")
DEFAULT_INCREMENTAL_OUTPUT_ROOT = Path("/media/jushen/project-rl-dataset/raw_data_0903_incremental")
UNKNOWN_STATION = "unknown_station"
UNKNOWN_TASK = "unknown_task"
SUCCESSFUL_DOWNLOAD_STATUSES = {"downloaded", "skipped_complete"}


@dataclass(frozen=True)
class DownloadItem:
    record_id: str
    excel_row: int
    sheet: str
    station_name_zh: str | None
    station_id: str
    task_name_zh: str | None
    task_id: str
    task_description_en: str | None
    task_description_reviewed: bool
    bos_path: str
    source_batch: str
    destination: str
    task_resolution_source: str = "legacy"
    task_source_row: int | None = None
    source_override_key: str | None = None


@dataclass(frozen=True)
class DownloadResult:
    record_id: str
    bos_path: str
    destination: str
    status: str
    hdf5_count: int
    total_bytes: int
    error: str | None
    completed_at: str


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def load_mapping(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload.get("stations"), dict) or not isinstance(payload.get("tasks"), dict):
        raise ValueError("mapping JSON requires object fields: stations, tasks")
    if not isinstance(payload.get("source_overrides", {}), dict):
        raise ValueError("mapping JSON field source_overrides must be an object")
    if not isinstance(payload.get("path_task_patterns", {}), dict):
        raise ValueError("mapping JSON field path_task_patterns must be an object")
    return payload


def find_named_columns(headers: Iterable[Any], required: list[str]) -> dict[str, int]:
    """Locate a caller-defined set of Excel columns."""
    names = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
    missing = [name for name in required if name not in names]
    if missing:
        raise ValueError(f"missing Excel columns: {missing}")
    return {name: names[name] for name in required}


def find_columns(headers: Iterable[Any]) -> dict[str, int]:
    required = ["任务描述", "机器编号", "采集状态", "原始数据存储路径", "父记录"]
    return find_named_columns(headers, required)


def extract_hyperlink(value: Any) -> str:
    """Extract BOS path from plain text or Excel HYPERLINK formula."""
    if value is None:
        return ""
    text = str(value).strip()
    match = re.search(
        r'HYPERLINK\(\s*"(?P<target>(?:[^"]|"")*)"\s*[,;]\s*"(?P<label>(?:[^"]|"")*)"\s*\)',
        text,
        flags=re.IGNORECASE,
    )
    if match:
        target = match.group("target").replace('""', '"').strip()
        label = match.group("label").replace('""', '"').strip()
        text = label if "BOS::" in label or label.startswith("bos:/") else target
    if text.startswith("http://BOS::"):
        text = text[len("http://") :]
    return text.strip()


def safe_component(value: str, fallback: str) -> str:
    text = value.strip().replace("/", "_").replace("\\", "_")
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fff]", "_", text)
    text = re.sub(r"_+", "_", text).strip("._")
    return text or fallback


def resolve_station(value: Any, mapping: dict[str, Any]) -> tuple[str | None, str]:
    name = str(value).strip() if value not in (None, "") else None
    if name is None:
        return None, UNKNOWN_STATION
    station_id = mapping["stations"].get(name)
    return name, safe_component(str(station_id), UNKNOWN_STATION) if station_id else UNKNOWN_STATION


def resolve_task(
    task_value: Any,
    parent_value: Any,
    inherited_task_name: str | None,
    mapping: dict[str, Any],
) -> tuple[str | None, str, str | None, bool, str]:
    if task_value not in (None, ""):
        task_name = str(task_value).strip()
        resolution_source = "direct"
    elif parent_value not in (None, ""):
        task_name = str(parent_value).strip()
        resolution_source = "parent"
    elif inherited_task_name is not None:
        task_name = inherited_task_name
        resolution_source = "inherited"
    else:
        task_name = None
        resolution_source = "unknown"
    entry = mapping["tasks"].get(task_name) if task_name else None
    if not isinstance(entry, dict):
        return task_name, UNKNOWN_TASK, None, False, resolution_source
    task_id = safe_component(str(entry.get("task_id", "")), UNKNOWN_TASK)
    description = entry.get("task_description_en")
    reviewed = bool(entry.get("reviewed", False))
    return task_name, task_id, str(description) if description else None, reviewed, resolution_source


def raw_cell_sha256(value: Any) -> str:
    return hashlib.sha256(str(value).encode("utf-8")).hexdigest()


def resolve_source_entries(
    *,
    worksheet_name: str,
    row_number: int,
    raw_value: Any,
    task_name: str | None,
    mapping: dict[str, Any],
) -> list[tuple[str, str, str | None]]:
    """Return (BOS path, destination batch name, override key) entries."""
    override_key = f"{worksheet_name}:{row_number}"
    override = mapping.get("source_overrides", {}).get(override_key)
    if override is None:
        path = extract_hyperlink(raw_value)
        return [(path, source_batch_name(path), None)]
    if not isinstance(override, dict):
        raise ValueError(f"source override {override_key} must be an object")

    expected_task = override.get("expected_task")
    if expected_task != task_name:
        raise ValueError(
            f"source override {override_key} task mismatch: expected={expected_task!r}, actual={task_name!r}"
        )
    expected_hash = override.get("expected_raw_sha256")
    actual_hash = raw_cell_sha256(raw_value)
    if expected_hash != actual_hash:
        raise ValueError(
            f"source override {override_key} raw cell changed: expected_sha256={expected_hash}, actual_sha256={actual_hash}"
        )

    sources = override.get("sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError(f"source override {override_key} requires a non-empty sources list")
    entries: list[tuple[str, str, str | None]] = []
    seen_paths: set[str] = set()
    seen_batches: set[str] = set()
    for index, source in enumerate(sources):
        if not isinstance(source, dict):
            raise ValueError(f"source override {override_key} sources[{index}] must be an object")
        path = str(source.get("bos_path", "")).strip()
        batch = safe_component(str(source.get("source_batch", "")), "")
        if path.count("BOS::") != 1 or not path.startswith("BOS::"):
            raise ValueError(f"source override {override_key} sources[{index}] has invalid BOS path")
        if not batch:
            raise ValueError(f"source override {override_key} sources[{index}] requires source_batch")
        if path in seen_paths:
            raise ValueError(f"source override {override_key} contains duplicate BOS path: {path}")
        if batch in seen_batches:
            raise ValueError(f"source override {override_key} contains duplicate source_batch: {batch}")
        seen_paths.add(path)
        seen_batches.add(batch)
        entries.append((path, batch, override_key))
    return entries


def bos_source(path: str, bos_base: str) -> str:
    text = path.strip()
    if text.startswith("bos:/"):
        return text
    if text.startswith("BOS::"):
        text = text[len("BOS::") :]
    return f"{bos_base.rstrip('/')}/{text.lstrip('/')}"


def source_batch_name(path: str) -> str:
    text = path.rstrip("/")
    if text.startswith("BOS::"):
        text = text[len("BOS::") :]
    return safe_component(text.rsplit("/", 1)[-1], "unknown_batch")


def make_record_id(sheet: str, row: int, path: str) -> str:
    digest = hashlib.sha256(f"{sheet}\0{row}\0{path}".encode("utf-8")).hexdigest()[:16]
    return f"row_{row}_{digest}"


def parse_board(excel: Path, mapping_path: Path, output_root: Path) -> list[DownloadItem]:
    mapping = load_mapping(mapping_path)
    # This board currently carries a stale worksheet dimension (A1:A1) in its
    # XML. openpyxl read-only iteration trusts that cache and would expose only
    # column A, while normal mode correctly discovers A:T.
    workbook = load_workbook(excel, data_only=False, read_only=False)
    records: list[DownloadItem] = []
    for worksheet in workbook.worksheets:
        current_task_name: str | None = None
        current_task_row: int | None = None
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            continue
        columns = find_columns(headers)
        for row_number, row in enumerate(rows, start=2):
            direct_task_value = row[columns["任务描述"]]
            if direct_task_value not in (None, ""):
                current_task_name = str(direct_task_value).strip()
                current_task_row = row_number
            status = row[columns["采集状态"]]
            raw_value = row[columns["原始数据存储路径"]]
            if str(status).strip() != "已完成" or raw_value in (None, ""):
                continue
            station_name, station_id = resolve_station(row[columns["机器编号"]], mapping)
            task_name, task_id, description, reviewed, resolution_source = resolve_task(
                direct_task_value,
                row[columns["父记录"]],
                current_task_name,
                mapping,
            )
            task_source_row = row_number if resolution_source in {"direct", "parent"} else current_task_row
            source_entries = resolve_source_entries(
                worksheet_name=worksheet.title,
                row_number=row_number,
                raw_value=raw_value,
                task_name=task_name,
                mapping=mapping,
            )
            for path, batch, override_key in source_entries:
                destination = output_root / station_id / task_id / batch
                records.append(
                    DownloadItem(
                        record_id=make_record_id(worksheet.title, row_number, path),
                        excel_row=row_number,
                        sheet=worksheet.title,
                        station_name_zh=station_name,
                        station_id=station_id,
                        task_name_zh=task_name,
                        task_id=task_id,
                        task_description_en=description,
                        task_description_reviewed=reviewed,
                        bos_path=path,
                        source_batch=batch,
                        destination=str(destination),
                        task_resolution_source=resolution_source,
                        task_source_row=task_source_row,
                        source_override_key=override_key,
                    )
                )
    return records


def task_from_bos_path(path: str, mapping: dict[str, Any]) -> tuple[str, str | None, str | None, bool]:
    """Resolve task metadata from an explicit, reviewed path-pattern mapping."""
    matches = [
        str(task_id)
        for pattern, task_id in mapping.get("path_task_patterns", {}).items()
        if str(pattern).lower() in path.lower()
    ]
    if len(matches) != 1:
        return UNKNOWN_TASK, None, None, False
    task_id = safe_component(matches[0], UNKNOWN_TASK)
    for task_name, entry in mapping["tasks"].items():
        if isinstance(entry, dict) and entry.get("task_id") == task_id:
            description = entry.get("task_description_en")
            return task_id, task_name, str(description) if description else None, bool(entry.get("reviewed", False))
    return task_id, None, None, False


def load_successful_bos_paths(results_path: Path) -> set[str]:
    """Read exact BOS roots that previously produced at least one valid HDF5 file."""
    successful: set[str] = set()
    if not results_path.is_file():
        raise FileNotFoundError(f"historical download results not found: {results_path}")
    with results_path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"invalid JSON at {results_path}:{line_number}: {error}") from error
            if record.get("status") not in SUCCESSFUL_DOWNLOAD_STATUSES:
                continue
            if int(record.get("hdf5_count") or 0) < 1:
                continue
            path = str(record.get("bos_path", "")).strip().rstrip("/")
            if path:
                successful.add(path)
    return successful


def parse_nine_device_sheet(
    excel: Path,
    mapping_path: Path,
    output_root: Path,
    successful_bos_paths: set[str],
    repair_bos_paths: set[str],
) -> tuple[list[DownloadItem], int]:
    """Build only missing downloads from the ``9台天轶设备`` inventory."""
    mapping = load_mapping(mapping_path)
    workbook = load_workbook(excel, data_only=True, read_only=True)
    records: list[DownloadItem] = []
    candidate_count = 0
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            continue
        columns = find_named_columns(headers, ["任务英文名称", "原始路径"])
        for row_number, row in enumerate(rows, start=2):
            raw_value = row[columns["原始路径"]]
            if raw_value in (None, ""):
                continue
            path = extract_hyperlink(raw_value)
            if not path.startswith(("BOS::", "bos:/")):
                path = f"BOS::{path.lstrip('/')}"
            path = path.rstrip("/")
            candidate_count += 1
            if path in successful_bos_paths:
                continue

            task_id, task_name, description, reviewed = task_from_bos_path(path, mapping)
            batch = source_batch_name(path)
            category = "repaired" if path in repair_bos_paths else "new"
            destination = output_root / category / task_id / batch
            records.append(
                DownloadItem(
                    record_id=make_record_id(worksheet.title, row_number, path),
                    excel_row=row_number,
                    sheet=worksheet.title,
                    station_name_zh=None,
                    station_id=UNKNOWN_STATION,
                    task_name_zh=task_name,
                    task_id=task_id,
                    task_description_en=description,
                    task_description_reviewed=reviewed,
                    bos_path=path,
                    source_batch=batch,
                    destination=str(destination),
                    task_resolution_source="path_pattern",
                    task_source_row=row_number,
                )
            )
    return records, candidate_count


def validate_plan(items: list[DownloadItem]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    unresolved: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    by_source: dict[str, list[DownloadItem]] = {}
    by_destination: dict[str, list[DownloadItem]] = {}
    for item in items:
        issues: list[str] = []
        if not item.bos_path:
            issues.append("missing_bos_path")
        elif item.bos_path.count("BOS::") > 1:
            issues.append("multiple_bos_paths_in_one_cell")
        elif not (item.bos_path.startswith("BOS::") or item.bos_path.startswith("bos:/")):
            issues.append("unsupported_bos_path")
        if item.station_id == UNKNOWN_STATION:
            issues.append("unknown_station")
        if item.task_id == UNKNOWN_TASK:
            issues.append("unknown_task")
        if item.task_description_en is None:
            issues.append("missing_english_description")
        if not item.task_description_reviewed:
            issues.append("english_description_not_reviewed")
        if issues:
            unresolved.append({**asdict(item), "issues": issues})
        by_source.setdefault(item.bos_path, []).append(item)
        by_destination.setdefault(item.destination, []).append(item)
    for kind, groups in (("bos_path", by_source), ("destination", by_destination)):
        for value, records in groups.items():
            if value and len(records) > 1:
                duplicates.append({"kind": kind, "value": value, "record_ids": [x.record_id for x in records]})
    return unresolved, duplicates


def write_jsonl_atomic(path: Path, records: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temp, path)


def write_csv(path: Path, records: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not records:
        path.write_text("", encoding="utf-8")
        return
    keys: list[str] = []
    for record in records:
        for key in record:
            if key not in keys:
                keys.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys, extrasaction="ignore")
        writer.writeheader()
        for record in records:
            flat = {key: json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value for key, value in record.items()}
            writer.writerow(flat)


def read_plan(path: Path) -> list[DownloadItem]:
    items: list[DownloadItem] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                items.append(DownloadItem(**json.loads(line)))
    return items


def verify_download(directory: Path) -> tuple[int, int, list[str]]:
    files = sorted(directory.rglob("trajectory.hdf5"))
    errors: list[str] = []
    total_bytes = 0
    for path in files:
        try:
            total_bytes += path.stat().st_size
            with h5py.File(path, "r") as h5_file:
                list(h5_file.keys())
        except (OSError, ValueError) as error:
            errors.append(f"{path}: {error}")
    if not files:
        errors.append("no trajectory.hdf5 found")
    return len(files), total_bytes, errors


def complete_marker_matches(destination: Path, item: DownloadItem) -> bool:
    marker = destination / ".download_complete.json"
    if not marker.is_file():
        return False
    try:
        payload = json.loads(marker.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    return payload.get("record_id") == item.record_id and payload.get("bos_path") == item.bos_path


def download_item(
    item: DownloadItem,
    *,
    bcecmd: Path,
    bos_base: str,
    state_dir: Path,
    resume: bool,
) -> DownloadResult:
    if item.bos_path.count("BOS::") > 1 or not (
        item.bos_path.startswith("BOS::") or item.bos_path.startswith("bos:/")
    ):
        return DownloadResult(
            item.record_id,
            item.bos_path,
            item.destination,
            "failed_plan_validation",
            0,
            0,
            "BOS path must contain exactly one source path",
            utc_now(),
        )
    destination = Path(item.destination)
    partial = destination.with_name(destination.name + ".partial")
    if destination.exists():
        if complete_marker_matches(destination, item):
            count, total_bytes, errors = verify_download(destination)
            status = "skipped_complete" if not errors else "failed_verification"
            return DownloadResult(item.record_id, item.bos_path, str(destination), status, count, total_bytes, "; ".join(errors) or None, utc_now())
        return DownloadResult(item.record_id, item.bos_path, str(destination), "failed", 0, 0, "destination exists without matching completion marker", utc_now())
    if partial.exists() and not resume:
        return DownloadResult(item.record_id, item.bos_path, str(destination), "failed", 0, 0, "partial directory exists; use --resume", utc_now())
    partial.mkdir(parents=True, exist_ok=True)
    logs_dir = state_dir / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    log_path = logs_dir / f"{item.record_id}.log"
    command = [str(bcecmd), "bos", "sync", bos_source(item.bos_path, bos_base), str(partial)]
    try:
        with log_path.open("a", encoding="utf-8") as log:
            log.write(f"\n[{utc_now()}] command={json.dumps(command, ensure_ascii=False)}\n")
            completed = subprocess.run(command, stdout=log, stderr=subprocess.STDOUT, check=False)
    except OSError as error:
        return DownloadResult(item.record_id, item.bos_path, str(destination), "failed", 0, 0, str(error), utc_now())
    if completed.returncode != 0:
        return DownloadResult(item.record_id, item.bos_path, str(destination), "failed", 0, 0, f"bcecmd exited with {completed.returncode}; log={log_path}", utc_now())
    count, total_bytes, errors = verify_download(partial)
    if errors:
        return DownloadResult(item.record_id, item.bos_path, str(destination), "failed_verification", count, total_bytes, "; ".join(errors), utc_now())
    marker_payload = {
        "record_id": item.record_id,
        "bos_path": item.bos_path,
        "station_name_zh": item.station_name_zh,
        "station_id": item.station_id,
        "task_name_zh": item.task_name_zh,
        "task_id": item.task_id,
        "task_description_en": item.task_description_en,
        "task_description_reviewed": item.task_description_reviewed,
        "hdf5_count": count,
        "total_bytes": total_bytes,
        "completed_at": utc_now(),
    }
    (partial / ".download_complete.json").write_text(
        json.dumps(marker_payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    os.rename(partial, destination)
    return DownloadResult(item.record_id, item.bos_path, str(destination), "downloaded", count, total_bytes, None, utc_now())


def append_jsonl(path: Path, record: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def run_plan(args: argparse.Namespace) -> int:
    items = parse_board(args.excel, args.mapping, args.output)
    unresolved, duplicates = validate_plan(items)
    args.state_dir.mkdir(parents=True, exist_ok=True)
    write_jsonl_atomic(args.state_dir / "download_plan.jsonl", (asdict(item) for item in items))
    write_csv(args.state_dir / "download_plan.csv", [asdict(item) for item in items])
    write_csv(args.state_dir / "unresolved_metadata.csv", unresolved)
    write_csv(args.state_dir / "duplicate_sources.csv", duplicates)
    print(f"download records: {len(items)}")
    print(f"unresolved records: {len(unresolved)}")
    print(f"duplicate groups: {len(duplicates)}")
    print(f"plan: {args.state_dir / 'download_plan.jsonl'}")
    return 0


def write_missing_bos_report(
    path: Path,
    *,
    excel: Path,
    history_results: Path,
    candidate_count: int,
    items: list[DownloadItem],
) -> None:
    repaired = [item for item in items if "/repaired/" in item.destination]
    new = [item for item in items if "/new/" in item.destination]
    lines = [
        "# Incremental BOS download plan",
        "",
        f"- Source Excel: `{excel}`",
        f"- Historical results: `{history_results}`",
        f"- Candidate BOS paths: {candidate_count}",
        f"- Missing paths: {len(items)}",
        f"- Repaired batches: {len(repaired)}",
        f"- New batches: {len(new)}",
        "",
    ]
    for title, records in (("Repaired", repaired), ("New", new)):
        lines.extend([f"## {title}", ""])
        if not records:
            lines.append("None.")
        else:
            for item in records:
                lines.append(f"- Excel row {item.excel_row}: `{item.bos_path}`")
                lines.append(f"  - destination: `{item.destination}`")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def run_incremental_plan(args: argparse.Namespace) -> int:
    successful = load_successful_bos_paths(args.history_results)
    repair_paths = {path.strip().rstrip("/") for path in args.repair_bos_path}
    items, candidate_count = parse_nine_device_sheet(
        args.excel,
        args.mapping,
        args.output,
        successful,
        repair_paths,
    )
    unresolved, duplicates = validate_plan(items)
    args.state_dir.mkdir(parents=True, exist_ok=True)
    write_jsonl_atomic(args.state_dir / "download_plan.jsonl", (asdict(item) for item in items))
    write_csv(args.state_dir / "download_plan.csv", [asdict(item) for item in items])
    write_csv(args.state_dir / "unresolved_metadata.csv", unresolved)
    write_csv(args.state_dir / "duplicate_sources.csv", duplicates)
    write_missing_bos_report(
        args.state_dir / "missing_bos.md",
        excel=args.excel,
        history_results=args.history_results,
        candidate_count=candidate_count,
        items=items,
    )
    repaired_count = sum("/repaired/" in item.destination for item in items)
    print(f"candidate BOS paths: {candidate_count}")
    print(f"already downloaded: {candidate_count - len(items)}")
    print(f"incremental records: {len(items)} (repaired={repaired_count}, new={len(items) - repaired_count})")
    print(f"unresolved metadata records: {len(unresolved)}")
    print(f"duplicate groups: {len(duplicates)}")
    print(f"plan: {args.state_dir / 'download_plan.jsonl'}")
    print(f"report: {args.state_dir / 'missing_bos.md'}")
    return 0


def run_download(args: argparse.Namespace) -> int:
    if not args.bcecmd.is_file():
        raise FileNotFoundError(f"bcecmd not found: {args.bcecmd}")
    items = read_plan(args.plan)
    results_path = args.state_dir / "download_results.jsonl"
    failed = 0
    for index, item in enumerate(items, start=1):
        print(f"[{index}/{len(items)}] {item.station_id}/{item.task_id}/{item.source_batch}")
        result = download_item(
            item,
            bcecmd=args.bcecmd,
            bos_base=args.bos_base,
            state_dir=args.state_dir,
            resume=args.resume,
        )
        append_jsonl(results_path, asdict(result))
        print(f"  {result.status}: hdf5={result.hdf5_count} bytes={result.total_bytes}")
        if result.error:
            print(f"  error: {result.error}", file=sys.stderr)
        if result.status.startswith("failed"):
            failed += 1
            if args.stop_on_error:
                break
    print(f"failed: {failed}")
    return 1 if failed else 0


def run_inventory_command(args: argparse.Namespace) -> int:
    from inventory_bos import default_state_dir, run_inventory

    if args.state_dir is None:
        args.state_dir = default_state_dir()
    return run_inventory(args)


def run_inventory_self_test() -> int:
    from inventory_bos import run_self_test as run_inventory_self_test

    return run_inventory_self_test()


def run_status(args: argparse.Namespace) -> int:
    items = read_plan(args.plan)
    counts = {"complete": 0, "partial": 0, "missing": 0, "conflict": 0}
    for item in items:
        destination = Path(item.destination)
        partial = destination.with_name(destination.name + ".partial")
        if destination.exists():
            key = "complete" if complete_marker_matches(destination, item) else "conflict"
        elif partial.exists():
            key = "partial"
        else:
            key = "missing"
        counts[key] += 1
    for key, value in counts.items():
        print(f"{key}: {value}")
    return 0


def run_self_test() -> int:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        anomalous_raw = (
            '==HYPERLINK("http://BOS::raw_data/robot_128/aBOS::raw_data/robot_128/b", '
            '"BOS::raw_data/robot_128/aBOS::raw_data/robot_128/b")'
        )
        mapping_payload = {
            "stations": {"天轶9": "tienyi_9"},
            "tasks": {
                "线缆连接": {
                    "task_id": "plug_cables",
                    "task_description_en": "Plug cables.",
                    "reviewed": True,
                },
                "插管": {
                    "task_id": "insert_hose",
                    "task_description_en": "Insert hose.",
                    "reviewed": True,
                },
            },
            "source_overrides": {
                "异常路径:2": {
                    "expected_task": "线缆连接",
                    "expected_raw_sha256": raw_cell_sha256(anomalous_raw),
                    "sources": [
                        {"bos_path": "BOS::raw_data/robot_128/a", "source_batch": "batch_a"},
                        {"bos_path": "BOS::raw_data/robot_128/b", "source_batch": "batch_b"},
                    ],
                }
            },
            "path_task_patterns": {"plug_cables": "plug_cables", "insert_hose": "insert_hose"},
        }
        mapping = root / "mapping.json"
        mapping.write_text(
            json.dumps(mapping_payload, ensure_ascii=False),
            encoding="utf-8",
        )
        book = Workbook()
        sheet = book.active
        sheet.title = "任务管理"
        sheet.append(["任务描述", "机器编号", "采集状态", "原始数据存储路径", "父记录"])
        sheet.append(["线缆连接", "天轶9", "已完成", '==HYPERLINK("http://BOS::raw_data/robot_128/batch", "BOS::raw_data/robot_128/batch")', None])
        sheet.append([None, None, "已完成", "BOS::raw_data/robot_128/unknown", None])
        sheet.append([None, "天轶9", "已完成", "BOS::raw_data/robot_128/parent", "线缆连接"])
        sheet.append(["插管", "天轶9", "已完成", "BOS::raw_data/robot_128/hose", None])
        sheet.append([None, "天轶9", "已完成", "BOS::raw_data/robot_128/hose_fail", None])

        first_blank = book.create_sheet("首行空任务")
        first_blank.append(["任务描述", "机器编号", "采集状态", "原始数据存储路径", "父记录"])
        first_blank.append([None, "天轶9", "已完成", "BOS::raw_data/robot_128/no_task", None])

        anomaly = book.create_sheet("异常路径")
        anomaly.append(["任务描述", "机器编号", "采集状态", "原始数据存储路径", "父记录"])
        anomaly.append(["线缆连接", "天轶9", "已完成", anomalous_raw, None])

        excel = root / "board.xlsx"
        book.save(excel)
        items = parse_board(excel, mapping, root / "output")
        assert len(items) == 8
        by_sheet_row = {(item.sheet, item.excel_row, item.bos_path): item for item in items}
        direct = by_sheet_row[("任务管理", 2, "BOS::raw_data/robot_128/batch")]
        inherited = by_sheet_row[("任务管理", 3, "BOS::raw_data/robot_128/unknown")]
        parent = by_sheet_row[("任务管理", 4, "BOS::raw_data/robot_128/parent")]
        switched = by_sheet_row[("任务管理", 6, "BOS::raw_data/robot_128/hose_fail")]
        unknown = by_sheet_row[("首行空任务", 2, "BOS::raw_data/robot_128/no_task")]
        override_items = [item for item in items if item.sheet == "异常路径"]
        assert direct.task_resolution_source == "direct" and direct.task_id == "plug_cables"
        assert inherited.task_resolution_source == "inherited" and inherited.task_source_row == 2
        assert inherited.station_id == UNKNOWN_STATION and "128" not in inherited.station_id
        assert parent.task_resolution_source == "parent" and parent.task_source_row == 4
        assert switched.task_resolution_source == "inherited" and switched.task_id == "insert_hose"
        assert unknown.task_resolution_source == "unknown" and unknown.task_id == UNKNOWN_TASK
        assert len(override_items) == 2
        assert {item.source_batch for item in override_items} == {"batch_a", "batch_b"}
        assert all(item.source_override_key == "异常路径:2" for item in override_items)

        mapping_without_override = root / "mapping_without_override.json"
        mapping_without_override.write_text(
            json.dumps({**mapping_payload, "source_overrides": {}}, ensure_ascii=False),
            encoding="utf-8",
        )
        unresolved_items = parse_board(excel, mapping_without_override, root / "output_without_override")
        unresolved, _ = validate_plan(unresolved_items)
        anomaly_issues = [
            record["issues"] for record in unresolved if record["sheet"] == "异常路径" and record["excel_row"] == 2
        ]
        assert anomaly_issues == [["multiple_bos_paths_in_one_cell"]]

        nine_device = Workbook()
        nine_sheet = nine_device.active
        nine_sheet.title = "Sheet1"
        nine_sheet.append(["任务英文名称", "原始路径"])
        nine_sheet.append(["plug", "raw_data/robot_128/plug_cables_downloaded"])
        nine_sheet.append(["plug", "raw_data/robot_128/plug_cables_new"])
        nine_sheet.append(["hose", "raw_data/robot_134/insert_hose_repair"])
        nine_excel = root / "nine.xlsx"
        nine_device.save(nine_excel)
        successful = {"BOS::raw_data/robot_128/plug_cables_downloaded"}
        repairs = {"BOS::raw_data/robot_134/insert_hose_repair"}
        incremental, candidate_count = parse_nine_device_sheet(
            nine_excel, mapping, root / "incremental", successful, repairs
        )
        assert candidate_count == 3 and len(incremental) == 2
        assert sum("/repaired/" in item.destination for item in incremental) == 1
        assert sum("/new/" in item.destination for item in incremental) == 1
    print("download_raw_data self-test: PASS")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    plan = subparsers.add_parser("plan", help="parse board and write a frozen download plan")
    plan.add_argument("--excel", type=Path, required=True)
    plan.add_argument("--mapping", type=Path, required=True)
    plan.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_ROOT)
    plan.add_argument("--state-dir", type=Path, required=True)
    plan.set_defaults(handler=run_plan)

    incremental = subparsers.add_parser(
        "plan-incremental",
        help="parse 9-device inventory and plan only BOS roots absent from historical successful results",
    )
    incremental.add_argument("--excel", type=Path, required=True)
    incremental.add_argument("--mapping", type=Path, required=True)
    incremental.add_argument("--history-results", type=Path, required=True)
    incremental.add_argument("--output", type=Path, default=DEFAULT_INCREMENTAL_OUTPUT_ROOT)
    incremental.add_argument("--state-dir", type=Path, required=True)
    incremental.add_argument(
        "--repair-bos-path",
        action="append",
        default=[],
        help="exact missing BOS root to place below output/repaired; repeat for multiple roots",
    )
    incremental.set_defaults(handler=run_incremental_plan)

    download = subparsers.add_parser("download", help="execute a frozen download plan")
    download.add_argument("--plan", type=Path, required=True)
    download.add_argument("--state-dir", type=Path, required=True)
    download.add_argument("--bcecmd", type=Path, default=DEFAULT_BCECMD)
    download.add_argument("--bos-base", default=DEFAULT_BOS_BASE)
    download.add_argument("--resume", action="store_true")
    download.add_argument("--stop-on-error", action="store_true")
    download.set_defaults(handler=run_download)

    status = subparsers.add_parser("status", help="show local status for a frozen plan")
    status.add_argument("--plan", type=Path, required=True)
    status.set_defaults(handler=run_status)

    inventory = subparsers.add_parser(
        "inventory",
        help="list BOS batch roots and compare them with local downloads; never syncs",
    )
    inventory.add_argument(
        "--excel",
        type=Path,
        default=Path(__file__).resolve().parent.parent.parent / "9台天轶设备.xlsx",
    )
    inventory.add_argument(
        "--mapping",
        type=Path,
        default=Path(__file__).resolve().parent / "raw_data_mapping.json",
    )
    inventory.add_argument("--bcecmd", type=Path, default=DEFAULT_BCECMD)
    inventory.add_argument("--bos-base", default=DEFAULT_BOS_BASE)
    inventory.add_argument("--local-root", type=Path, action="append")
    inventory.add_argument("--results", type=Path, action="append")
    inventory.add_argument("--state-dir", type=Path)
    inventory.add_argument("--skip-bos-list", action="store_true")
    inventory.set_defaults(handler=run_inventory_command, list_children=None)

    inventory_test = subparsers.add_parser("inventory-self-test")
    inventory_test.set_defaults(handler=lambda _: run_inventory_self_test())

    self_test = subparsers.add_parser("self-test")
    self_test.set_defaults(handler=lambda _: run_self_test())
    return parser


def main() -> None:
    args = build_parser().parse_args()
    raise SystemExit(args.handler(args))


if __name__ == "__main__":
    main()
