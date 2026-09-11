#!/usr/bin/env python3
"""Generate the formatted production internship report DOCX from Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PRODUCTION_INTERNSHIP_REPORT.md"
OUTPUT = ROOT / "PRODUCTION_INTERNSHIP_REPORT.docx"

FIGURES = {
    "production": ROOT / "cosmos_production_internship_ppt/origin_image/slide_03.png",
    "workflow": ROOT / "cosmos_production_internship_ppt/origin_image/slide_04.png",
    "architecture": ROOT / "cosmos_production_internship_ppt/origin_image/slide_06.png",
    "translation": ROOT / "figures/compare/fig2_translation_2x2_report.png",
    "psnr": ROOT / "figures/compare/fig8_future_cameras_psnr.png",
    "spearman": ROOT / "figures/compare/fig9_value_spearman.png",
    "rgb": ROOT / "figures/rgb_examples/rotation6d_20d_policy_2000_0804_am_success_0804_095314_row750.png",
}

BLUE = "003F8F"
LIGHT_BLUE = "EAF2FF"
LIGHT_GRAY = "F3F6FA"
RED = "B5121B"


def set_east_asia_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str, placeholder: str = ""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_specs = {
        "Title": ("黑体", 22, WD_ALIGN_PARAGRAPH.CENTER, 0, 12),
        "Heading 1": ("黑体", 16, WD_ALIGN_PARAGRAPH.LEFT, 14, 7),
        "Heading 2": ("黑体", 14, WD_ALIGN_PARAGRAPH.LEFT, 10, 5),
        "Heading 3": ("黑体", 12, WD_ALIGN_PARAGRAPH.LEFT, 7, 3),
    }
    for style_name, (font, size, align, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True
        # 报告各级标题统一使用黑体、黑色，避免模板色干扰正式排版。
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = align
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)
        # python-docx 的默认 Title 样式带蓝色底边；正式报告标题不保留装饰线。
        p_pr = style._element.get_or_add_pPr()
        p_borders = p_pr.find(qn("w:pBdr"))
        if p_borders is not None:
            p_pr.remove(p_borders)

    if "图表题注" not in [s.name for s in doc.styles]:
        caption = doc.styles.add_style("图表题注", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["图表题注"]
    caption.font.name = "宋体"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.first_line_indent = Pt(0)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_header_footer(section):
    section.different_first_page_header_footer = True
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run("《工业工程生产实践》总结报告")
    set_east_asia_font(run, "宋体", 9)
    run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    add_field(p, "PAGE", "1")
    for run in p.runs:
        set_east_asia_font(run, "宋体", 9)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(32)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("《工业工程生产实践》\n生产实习总结报告")
    set_east_asia_font(run, "黑体", 22, True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(50)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("面向具身智能的机器人数据治理与\nCosmos Policy 离线中训练实践")
    set_east_asia_font(run, "黑体", 20, True)
    run.font.color.rgb = RGBColor(0, 0, 0)

    rows = [
        ("姓　　名", "葛铭博"),
        ("学　　号", "2023010345"),
        ("实习单位", "北京人形机器人创新中心有限公司"),
        ("实习部门", "具身智能部"),
        ("企业导师", "Chris Ren、Linda Zhao"),
        ("实习时间", "2026年7月1日—9月4日（10周）"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(11.0)
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(3.0 if idx == 0 else 11.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=130, bottom=130)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Pt(0)
                for run in paragraph.runs:
                    set_east_asia_font(run, "宋体", 11 if label == "实习时间" and idx == 1 else 12, idx == 0)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:val"), "nil")
                borders.append(element)
            tc_pr.append(borders)
    doc.add_page_break()


def add_toc(doc: Document):
    p = doc.add_paragraph("目录", style="Title")
    p.paragraph_format.space_after = Pt(18)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Pt(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "请在Word中右键更新目录")
    doc.add_page_break()


def add_inline(paragraph, text: str):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("\\sigma", "σ").replace("\\times", "×")
    # 避免中文技术词在跨页处被拆成两个部分。
    text = text.replace("优化器", "优\u2060化\u2060器")
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_east_asia_font(run, "宋体", 12, True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_east_asia_font(run, "Consolas", 10.5)
        else:
            run = paragraph.add_run(token)
            set_east_asia_font(run, "宋体", 12)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(style="图表题注")
    p.add_run(text)


def add_picture(doc: Document, path: Path, caption: str, width_cm: float = 14.6):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_two_pictures(doc: Document, left: Path, right: Path, caption: str):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, path in zip(table.rows[0].cells, (left, right)):
        cell.width = Cm(7.25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.add_run().add_picture(str(path), width=Cm(7.05))
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            borders.append(element)
        tc_pr.append(borders)
    add_caption(doc, caption)


def add_markdown_table(doc: Document, lines: list[str], table_index: int):
    parsed = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    if len(parsed) > 1 and all(re.fullmatch(r":?-{3,}:?", part.replace(" ", "")) for part in parsed[1]):
        parsed.pop(1)
    cols = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(parsed):
        for c_idx in range(cols):
            value = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = re.sub(r"\*\*|`", "", value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_east_asia_font(run, "宋体", 9, r_idx == 0)
                    if r_idx == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
    return table_index + 1


def add_code_block(doc: Document, code: list[str]):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    run = p.add_run("\n".join(code))
    set_east_asia_font(run, "等线", 9.5)


def add_equation(doc: Document, text: str):
    clean = text.replace("\\[", "").replace("\\]", "").strip()
    if "L_{EDM}" in clean:
        clean = "L_EDM = Σ[M · w(σ) · e²] / ΣM"
    else:
        clean = clean.replace("\\mathrm", "").replace("\\frac", "frac")
        clean = clean.replace("\\sigma", "σ").replace("\\times", "×").replace("\\sum", "Σ")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(clean)
    set_east_asia_font(run, "Cambria Math", 11)


def add_body(doc: Document):
    raw_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    # The Markdown title and metadata are represented by the dedicated cover.
    start = next(i for i, line in enumerate(raw_lines) if line.startswith("## 摘要"))
    lines = raw_lines[start:]
    i = 0
    in_code = False
    code: list[str] = []
    in_equation = False
    equation: list[str] = []
    table_index = 1
    inserted = set()

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code)
                code = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue

        if line.strip() == "\\[":
            in_equation = True
            equation = []
            i += 1
            continue
        if in_equation:
            if line.strip() == "\\]":
                add_equation(doc, " ".join(equation))
                in_equation = False
            else:
                equation.append(line.strip())
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            table_index = add_markdown_table(doc, table_lines, table_index)
            continue

        if not line.strip():
            i += 1
            continue

        if line.startswith("#### "):
            title = line[5:].strip()
            doc.add_paragraph(title, style="Heading 3")
            if title.startswith("4.3 ") and "architecture" not in inserted:
                add_picture(doc, FIGURES["architecture"], "图3　Cosmos Policy离线训练架构", 14.6)
                inserted.add("architecture")
            i += 1
            continue
        if line.startswith("### "):
            title = line[4:].strip()
            doc.add_paragraph(title, style="Heading 2")
            if title.startswith("4. 解决方案") and "workflow" not in inserted:
                add_picture(doc, FIGURES["workflow"], "图2　从原始机器人数据到可验证模型的总体改善方案", 14.6)
                inserted.add("workflow")
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            if title == "摘要":
                p = doc.add_paragraph("摘　要", style="Title")
                p.paragraph_format.space_after = Pt(8)
            else:
                doc.add_paragraph(title, style="Heading 1")
            i += 1
            continue

        if line.startswith("# "):
            i += 1
            continue

        if re.match(r"^[-*] ", line):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.add_run("•　")
            add_inline(p, re.sub(r"^[-*] ", "", line))
            i += 1
            continue

        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            match = re.match(r"^(\d+)\. (.*)$", line)
            p.add_run(f"{match.group(1)}.　")
            add_inline(p, match.group(2))
            i += 1
            continue

        if line.startswith("**关键词：**"):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline(p, line)
            doc.add_page_break()
            i += 1
            continue

        p = doc.add_paragraph(style="Normal")
        add_inline(p, line)

        if line.startswith("本项目将运作流程整理为两个相互衔接的阶段") and "production" not in inserted:
            add_picture(doc, FIGURES["production"], "图1　真实多视角机器人数据来源与生产流程痛点", 14.6)
            inserted.add("production")

        if line.startswith("结果表明："):
            pass

        if line.startswith("注：translation MAE") and "translation" not in inserted:
            add_picture(doc, FIGURES["translation"], "图4　Euler/rotation-6D与Policy/Joint的平移误差对比", 13.4)
            inserted.add("translation")

        if line.startswith("E-P 与 6D-P 到 2,000 steps") and "world_value" not in inserted:
            add_two_pictures(doc, FIGURES["psnr"], FIGURES["spearman"], "图5　未来相机PSNR与Value Spearman对比")
            add_picture(doc, FIGURES["rgb"], "图6　6D-P未来RGB预测示例（GT｜Prediction｜Absolute Error）", 14.6)
            inserted.add("world_value")

        i += 1


def main():
    for name, path in FIGURES.items():
        if not path.exists():
            raise FileNotFoundError(f"Missing figure {name}: {path}")

    doc = Document()
    configure_document(doc)
    configure_header_footer(doc.sections[0])
    add_cover(doc)
    add_toc(doc)
    add_body(doc)

    core = doc.core_properties
    core.title = "面向具身智能的机器人数据治理与Cosmos Policy离线中训练实践"
    core.subject = "《工业工程生产实践》总结报告"
    core.author = "葛铭博"
    core.keywords = "具身智能, 机器人数据治理, Cosmos Policy, 离线训练, 分布式训练"

    doc.save(OUTPUT)
    print(f"saved={OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")


if __name__ == "__main__":
    main()
